from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
BASE=Path('data')
bottom=pd.read_csv(BASE/'bottom_dataset.csv')
central=Path('results')/'Bottom_Datasets_All_Results.xlsx'
if not central.exists():
    print('Central missing')
    raise SystemExit
wb=load_workbook(central, read_only=True)
sheets=set(wb.sheetnames)
missing_in_central=[]
missing_files=[]
for ds in bottom['dataset']:
    name=ds.replace('.tsv.gz','')
    invalid=set('[]:*?/\\')
    s=''.join(c for c in name if c not in invalid)
    if len(s)>31: s=s[:31]
    # check per-dataset filled file exists
    per=BASE/'Bottom_Datasets_mlp'/name/'MLP Evaluation sheet(final)_filled.xlsx'
    if not per.exists():
        missing_files.append(ds)
    if s not in sheets:
        missing_in_central.append(ds)
print('Total bottom datasets:', len(bottom))
print('Per-dataset filled files missing:', len(missing_files))
for m in missing_files:
    print('  MISSING FILE:', m)
print('\nDatasets not present in central workbook:', len(missing_in_central))
for m in missing_in_central:
    print('  NOT IN CENTRAL:', m)
