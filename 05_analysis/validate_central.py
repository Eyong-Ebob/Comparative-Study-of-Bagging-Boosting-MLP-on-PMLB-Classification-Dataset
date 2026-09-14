from pathlib import Path
import sys
from openpyxl import load_workbook
import pandas as pd

central = Path('results') / 'Bottom_Datasets_All_Results.xlsx'
print('Central exists:', central.exists())
if not central.exists():
    sys.exit(0)

try:
    wb = load_workbook(central, read_only=True)
    sheets = wb.sheetnames
    print('Sheets found:', len(sheets))
    for s in sheets:
        print('\nSheet:', s)
        try:
            df = pd.read_excel(central, sheet_name=s)
            print('  shape:', df.shape)
            print('  columns:', list(df.columns)[:10])
            print('  first row:', df.head(1).to_dict(orient='records'))
        except Exception as e:
            print('  Failed reading sheet with pandas:', e)
except Exception as e:
    print('Failed to open workbook with openpyxl:', type(e).__name__, e)
