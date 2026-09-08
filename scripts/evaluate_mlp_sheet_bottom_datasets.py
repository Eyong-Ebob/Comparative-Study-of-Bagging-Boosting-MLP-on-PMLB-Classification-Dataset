from __future__ import annotations

import argparse
import json
import random
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from openpyxl import load_workbook
from sklearn.metrics import accuracy_score, balanced_accuracy_score, f1_score
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import LabelEncoder, StandardScaler
from torch.utils.data import DataLoader, TensorDataset

ROOT = Path(__file__).resolve().parent
TEMPLATE_XLSX = ROOT / "results" / "MLP Evaluation sheet(final).xlsx"
BOTTOM_ROOT = ROOT / "data" / "Bottom_Datasets_mlp"
DATA_ROOT = ROOT / "data"

RANDOM_SEED = 90483257
TEST_SIZE = 0.20


@dataclass
class ParamRow:
    row_idx: int
    hidden_neurons: int
    activation: str
    learning_rate: float
    batch_size: int
    epochs: int


class TabularMLP(nn.Module):
    def __init__(self, in_features: int, hidden_neurons: int, out_features: int, activation: str):
        super().__init__()
        act = self._get_activation(activation)
        self.net = nn.Sequential(
            nn.Linear(in_features, hidden_neurons),
            act,
            nn.Linear(hidden_neurons, hidden_neurons),
            act,
            nn.Linear(hidden_neurons, out_features),
        )

    @staticmethod
    def _get_activation(name: str) -> nn.Module:
        key = str(name).strip().lower()
        if key == "relu":
            return nn.ReLU()
        if key == "leaky_relu":
            return nn.LeakyReLU(negative_slope=0.01)
        if key == "gelu":
            return nn.GELU()
        if key == "tanh":
            return nn.Tanh()
        if key == "sigmoid":
            return nn.Sigmoid()
        raise ValueError(f"Unsupported activation: {name}")

    def forward(self, x):
        return self.net(x)


def set_seed(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed(seed)
        torch.cuda.manual_seed_all(seed)
    torch.backends.cudnn.deterministic = True
    torch.backends.cudnn.benchmark = False


def parse_param_rows(template_path: Path) -> list[ParamRow]:
    wb = load_workbook(template_path, data_only=True)
    ws = wb["Sheet1"]

    rows: list[ParamRow] = []
    for r in range(1, ws.max_row + 1):
        hidden = ws.cell(r, 2).value
        activation = ws.cell(r, 4).value
        lr = ws.cell(r, 5).value
        batch = ws.cell(r, 6).value
        epochs = ws.cell(r, 7).value

        if hidden is None or activation is None or lr is None or batch is None or epochs is None:
            continue

        if not isinstance(hidden, (int, float)):
            continue

        rows.append(
            ParamRow(
                row_idx=r,
                hidden_neurons=int(hidden),
                activation=str(activation).strip(),
                learning_rate=float(lr),
                batch_size=int(batch),
                epochs=int(epochs),
            )
        )

    if not rows:
        raise RuntimeError("No parameter rows were found in the template workbook.")
    return rows


def resolve_dataset_file(dataset_dir: Path) -> Path:
    ds_name = dataset_dir.name
    local = dataset_dir / f"{ds_name}.tsv.gz"
    if local.exists():
        return local

    fallback = DATA_ROOT / f"{ds_name}.tsv.gz"
    if fallback.exists():
        return fallback

    matches = list(DATA_ROOT.glob(f"{ds_name}.tsv"))
    if matches:
        return matches[0]

    raise FileNotFoundError(f"Dataset file not found for {ds_name}")


def load_dataset(dataset_path: Path) -> tuple[np.ndarray, np.ndarray]:
    df = pd.read_csv(dataset_path, sep="\t", compression="infer")

    if "class" in df.columns:
        target_col = "class"
    elif "target" in df.columns:
        target_col = "target"
    else:
        target_col = df.columns[-1]

    X = df.drop(columns=[target_col]).astype(np.float32).values
    y_raw = df[target_col].values

    encoder = LabelEncoder()
    y = encoder.fit_transform(y_raw)
    return X, y


def run_one_config(X: np.ndarray, y: np.ndarray, config: ParamRow, device: torch.device) -> tuple[float, float, float, float]:
    set_seed(RANDOM_SEED)

    X_train, X_test, y_train, y_test = train_test_split(
        X,
        y,
        test_size=TEST_SIZE,
        random_state=RANDOM_SEED,
        stratify=y,
    )

    scaler = StandardScaler()
    X_train = scaler.fit_transform(X_train).astype(np.float32)
    X_test = scaler.transform(X_test).astype(np.float32)

    num_classes = len(np.unique(y))
    binary = num_classes == 2
    out_dim = 1 if binary else num_classes

    X_train_t = torch.tensor(X_train, dtype=torch.float32, device=device)
    X_test_t = torch.tensor(X_test, dtype=torch.float32, device=device)

    if binary:
        y_train_t = torch.tensor(y_train.astype(np.float32), dtype=torch.float32, device=device).view(-1, 1)
    else:
        y_train_t = torch.tensor(y_train.astype(np.int64), dtype=torch.long, device=device)

    train_loader = DataLoader(
        TensorDataset(X_train_t, y_train_t),
        batch_size=max(1, config.batch_size),
        shuffle=True,
    )

    model = TabularMLP(X.shape[1], config.hidden_neurons, out_dim, config.activation).to(device)
    criterion = nn.BCEWithLogitsLoss() if binary else nn.CrossEntropyLoss()
    optimizer = torch.optim.Adam(model.parameters(), lr=config.learning_rate)

    model.train()
    start = time.perf_counter()
    for _ in range(config.epochs):
        for batch_X, batch_y in train_loader:
            optimizer.zero_grad()
            logits = model(batch_X)
            loss = criterion(logits, batch_y)
            loss.backward()
            optimizer.step()
    elapsed = time.perf_counter() - start

    model.eval()
    with torch.no_grad():
        logits = model(X_test_t)
        if binary:
            probs = torch.sigmoid(logits).squeeze(1).cpu().numpy()
            y_pred = (probs >= 0.5).astype(int)
        else:
            probs = torch.softmax(logits, dim=1).cpu().numpy()
            y_pred = np.argmax(probs, axis=1).astype(int)

    bal_acc = balanced_accuracy_score(y_test, y_pred)
    acc = accuracy_score(y_test, y_pred)
    f1 = f1_score(y_test, y_pred, average="macro", zero_division=0)
    return bal_acc, acc, f1, elapsed


def iter_dataset_dirs(limit: int | None = None) -> Iterable[Path]:
    dirs = sorted(d for d in BOTTOM_ROOT.iterdir() if d.is_dir())
    if limit is not None:
        dirs = dirs[:limit]
    return dirs


def evaluate_all(output_name: str, limit: int | None = None, resume: bool = True, dataset_name: str | None = None) -> None:
    if not TEMPLATE_XLSX.exists():
        raise FileNotFoundError(f"Template workbook not found: {TEMPLATE_XLSX}")

    param_rows = parse_param_rows(TEMPLATE_XLSX)
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

    manifest = []
    failures = []

    for ds_dir in iter_dataset_dirs(limit=limit):
        if dataset_name is not None and ds_dir.name != dataset_name:
            continue
        ds_name = ds_dir.name
        print(f"\n=== Dataset: {ds_name} ===")

        out_xlsx = ds_dir / output_name
        out_csv = ds_dir / "mlp_evaluation_sheet_results.csv"
        if resume and out_xlsx.exists() and out_csv.exists():
            print("  Skipped (resume: output files already exist)")
            manifest.append(
                {
                    "dataset": ds_name,
                    "dataset_file": None,
                    "output_xlsx": str(out_xlsx),
                    "output_csv": str(out_csv),
                    "rows_filled": None,
                    "skipped": True,
                }
            )
            continue

        wb = load_workbook(TEMPLATE_XLSX)
        ws = wb["Sheet1"]

        try:
            dataset_path = resolve_dataset_file(ds_dir)
            X, y = load_dataset(dataset_path)

            combo_results = []
            for idx, prm in enumerate(param_rows, start=1):
                print(
                    f"  [{idx}/{len(param_rows)}] row={prm.row_idx} h={prm.hidden_neurons} act={prm.activation} "
                    f"lr={prm.learning_rate} bs={prm.batch_size} ep={prm.epochs}",
                    flush=True,
                )
                bal_acc, acc, f1, elapsed = run_one_config(X, y, prm, device)
                ws.cell(prm.row_idx, 8).value = float(bal_acc)
                ws.cell(prm.row_idx, 9).value = float(acc)
                ws.cell(prm.row_idx, 10).value = float(f1)
                ws.cell(prm.row_idx, 11).value = float(acc)
                # record elapsed time (seconds) in column 12
                try:
                    ws.cell(prm.row_idx, 12).value = float(elapsed)
                except Exception:
                    ws.cell(prm.row_idx, 12).value = None

                combo_results.append(
                    {
                        "row": prm.row_idx,
                        "hidden_neurons": prm.hidden_neurons,
                        "activation": prm.activation,
                        "learning_rate": prm.learning_rate,
                        "batch_size": prm.batch_size,
                        "epochs": prm.epochs,
                        "balanced_accuracy": bal_acc,
                        "accuracy": acc,
                        "f1_score": f1,
                        "Time(s)": elapsed,
                    }
                )

            wb.save(out_xlsx)
            pd.DataFrame(combo_results).to_csv(out_csv, index=False)

            manifest.append(
                {
                    "dataset": ds_name,
                    "dataset_file": str(dataset_path),
                    "output_xlsx": str(out_xlsx),
                    "output_csv": str(out_csv),
                    "rows_filled": len(combo_results),
                }
            )

        except Exception as exc:
            failures.append({"dataset": ds_name, "error": str(exc)})
            wb.save(out_xlsx)
            print(f"  FAILED: {exc}")

    (ROOT / "results" / "mlp_evaluation_sheet_manifest.json").write_text(
        json.dumps({"completed": manifest, "failures": failures}, indent=2),
        encoding="utf-8",
    )

    print("\nDone.")
    print(f"Completed datasets: {len(manifest)}")
    print(f"Failed datasets: {len(failures)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output-name",
        default="MLP Evaluation sheet(final)_filled.xlsx",
        help="Filename for the copied and filled workbook in each dataset folder.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Optional number of datasets to process for testing.",
    )
    parser.add_argument(
        "--no-resume",
        action="store_true",
        help="Recompute datasets even if output files already exist.",
    )
    parser.add_argument(
        "--dataset",
        type=str,
        default=None,
        help="Optional single dataset name to process (folder name under data/Bottom_Datasets_mlp).",
    )
    args = parser.parse_args()

    evaluate_all(output_name=args.output_name, limit=args.limit, resume=not args.no_resume, dataset_name=args.dataset)


if __name__ == "__main__":
    main()
