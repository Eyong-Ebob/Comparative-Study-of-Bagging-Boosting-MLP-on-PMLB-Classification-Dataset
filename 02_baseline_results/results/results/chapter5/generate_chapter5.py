import warnings
warnings.filterwarnings('ignore')
import pandas as pd
import numpy as np
from pathlib import Path
import matplotlib.pyplot as plt
import seaborn as sns
from scipy import stats
from sklearn.linear_model import LinearRegression
from docx import Document
from docx.shared import Inches

ROOT = Path('.')
OUT = ROOT/'results'/'chapter5'
OUT.mkdir(parents=True, exist_ok=True)
FIGS = OUT/'figs'
FIGS.mkdir(parents=True, exist_ok=True)

def load_master():
    f = ROOT/'results'/'analysis'/'master_results_fixed.csv'
    if not f.exists():
        raise FileNotFoundError(f'Master results not found: {f}')
    df = pd.read_csv(f, index_col=0)
    return df

def descriptive_stats(df, suffix=''):
    desc = df.describe().T
    desc = desc[['count','mean','std','50%','min','max']]
    desc = desc.rename(columns={'50%':'median'})
    out = OUT/f'descriptive_stats_ch5{suffix}.csv'
    desc.to_csv(out)
    return desc

def ba_boxplot(df, suffix=''):
    data = df[['LogisticRegression','RandomForest','XGB','MLP']]
    data_m = data.melt(var_name='Classifier', value_name='Balanced_Accuracy')
    plt.figure(figsize=(8,6))
    sns.boxplot(x='Classifier', y='Balanced_Accuracy', data=data_m)
    plt.title('Balanced Accuracy by Classifier')
    plt.tight_layout()
    p = FIGS/f'ba_boxplot{suffix}.png'
    plt.savefig(p)
    plt.close()
    return p

def mean_ranks(df, suffix=''):
    ranks = df[['LogisticRegression','RandomForest','XGB','MLP']].rank(axis=1, method='average', ascending=False)
    meanr = ranks.mean().to_frame('mean_rank')
    out = OUT/f'mean_ranks_ch5{suffix}.csv'
    meanr.to_csv(out)
    return meanr

def win_loss_tie(df, suffix=''):
    cols = ['LogisticRegression','RandomForest','XGB','MLP']
    base = df['MLP']
    results = {}
    for c in cols[:-1]:
        comp = df[c]
        win = (base>comp).sum()
        loss = (base<comp).sum()
        tie = (base==comp).sum()
        results[c]=dict(win=win,loss=loss,tie=tie)
    out = pd.DataFrame(results).T
    out_file = OUT/f'win_loss_tie_MLP_vs_others{suffix}.csv'
    out.to_csv(out_file)
    return out

def wilcoxon_tests(df, suffix=''):
    cols = ['LogisticRegression','RandomForest','XGB','MLP']
    pairs = []
    rows = []
    for i in range(len(cols)):
        for j in range(i+1,len(cols)):
            a = df[cols[i]].dropna()
            b = df[cols[j]].dropna()
            idx = df[[cols[i],cols[j]]].dropna().index
            if len(idx)<5:
                pval = np.nan
                stat = np.nan
            else:
                stat, pval = stats.wilcoxon(df.loc[idx,cols[i]], df.loc[idx,cols[j]])
            rows.append({'pair':f'{cols[i]}_vs_{cols[j]}','wilcoxon_stat':stat,'p_value':pval,'n':len(idx)})
    out = pd.DataFrame(rows).set_index('pair')
    out_file = OUT/f'wilcoxon_ch5{suffix}.csv'
    out.to_csv(out_file)
    return out

def spearman_correlations(df, suffix=''):
    cols = ['LogisticRegression','RandomForest','XGB','MLP']
    corr = df[cols].corr(method='spearman')
    out_file = OUT/f'spearman_correlations_ch5{suffix}.csv'
    corr.to_csv(out_file)
    return corr

def regression_mlp_on_others(df, suffix=''):
    joint = df[['RandomForest','XGB','MLP']].dropna()
    if len(joint) < 5:
        res = {'r2':np.nan, 'intercept':np.nan, 'RandomForest':np.nan, 'XGB':np.nan}
        return res, None
    X = joint[['RandomForest','XGB']]
    y = joint['MLP']
    model = LinearRegression()
    model.fit(X, y)
    preds = model.predict(X)
    r2 = model.score(X,y)
    coef = dict(zip(X.columns, model.coef_))
    res = {'r2':r2, 'intercept':model.intercept_}
    res.update(coef)
    pd.Series(res).to_frame('value').to_csv(OUT/f'regression_mlp_on_rf_xgb{suffix}.csv')
    plt.figure(figsize=(6,5))
    plt.scatter(preds, y, alpha=0.6)
    plt.xlabel('Predicted MLP BA')
    plt.ylabel('Actual MLP BA')
    plt.title('MLP BA: Predicted vs Actual (RF+XGB)')
    plt.tight_layout()
    p = FIGS/f'regression_pred_actual{suffix}.png'
    plt.savefig(p)
    plt.close()
    return res, p


def create_doc(summary_tables, figs, suffix=''):
    doc = Document()
    doc.add_heading('Chapter 5 — Experimental Results and Analysis', level=1)
    doc.add_paragraph('This document summarizes the comparative results across classifiers and the MLP hyperparameter search analyses.')
    doc.add_heading('Descriptive Statistics', level=2)
    doc.add_paragraph(f'See attached CSV: descriptive_stats_ch5{suffix}.csv')
    doc.add_heading('Balanced Accuracy Boxplot', level=2)
    doc.add_picture(str(figs['ba_boxplot']), width=Inches(6))
    doc.add_heading('Mean Ranks', level=2)
    doc.add_paragraph(f'See attached CSV: mean_ranks_ch5{suffix}.csv')
    doc.add_heading('Wilcoxon Paired Tests', level=2)
    doc.add_paragraph(f'See attached CSV: wilcoxon_ch5{suffix}.csv')
    doc.add_heading('Spearman Correlations', level=2)
    doc.add_paragraph(f'See attached CSV: spearman_correlations_ch5{suffix}.csv')
    doc.add_heading('Regression (MLP ~ RF + XGB)', level=2)
    if figs.get('regression_pred_actual') is not None:
        doc.add_picture(str(figs['regression_pred_actual']), width=Inches(5))
    outdoc = OUT/f'Chapter5_Report{suffix}.docx'
    doc.save(outdoc)
    return outdoc


def save_filtered_master(df, suffix=''):
    out = OUT/f'master_results_mlp_only{suffix}.csv'
    df.to_csv(out)
    return out

def create_filtered_analysis(df, suffix=''):
    save_filtered_master(df, suffix)
    desc = descriptive_stats(df, suffix)
    p1 = ba_boxplot(df, suffix)
    meanr = mean_ranks(df, suffix)
    wlt = win_loss_tie(df, suffix)
    wil = wilcoxon_tests(df, suffix)
    spe = spearman_correlations(df, suffix)
    reg_res, p2 = regression_mlp_on_others(df, suffix)
    figs = {'ba_boxplot':p1, 'regression_pred_actual':p2}
    doc = create_doc(None, figs, suffix)
    return doc

def main():
    df = load_master()
    # Full analysis including all master rows
    desc = descriptive_stats(df)
    p1 = ba_boxplot(df)
    meanr = mean_ranks(df)
    wlt = win_loss_tie(df)
    wil = wilcoxon_tests(df)
    spe = spearman_correlations(df)
    reg_res, p2 = regression_mlp_on_others(df)
    figs = {'ba_boxplot':p1, 'regression_pred_actual':p2}
    doc = create_doc(None, figs)
    print('Wrote outputs to', OUT)
    print('Doc:', doc)
    # Filtered analysis: only datasets with successful MLP results
    mlp_only = df[df['MLP'].notna()]
    doc_filtered = create_filtered_analysis(mlp_only, '_MLP_Successful')
    print('Wrote filtered outputs to', OUT)
    print('Filtered Doc:', doc_filtered)

if __name__=='__main__':
    main()
